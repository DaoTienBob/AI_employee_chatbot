"""Document text extraction (T05).

Extracts readable text from PDF (pypdf) and DOCX (python-docx) files,
returning sections (heading + body paragraphs) so downstream chunking can
attach usable section metadata (roadmap §3.2: section=Employee Benefits).
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

# Headings-ish lines: short, no terminal punctuation, optionally numbered
# ("1. Purpose", "2.1 Scope", "Section 3 - Leave").
_HEADING_PATTERN = re.compile(
    r"^(?:(?:section\s+)?\d+(?:\.\d+)*[.):\-]?\s+)?[A-Z][A-Za-z0-9 ,&/()'\-]{2,80}$"
)
_MAX_HEADING_WORDS = 12


@dataclass
class Section:
    """A contiguous block of text under one heading (or one PDF page)."""

    title: str
    paragraphs: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p for p in self.paragraphs if p.strip())


@dataclass
class ExtractedDocument:
    """Full extraction result for one file."""

    sections: list[Section] = field(default_factory=list)

    @property
    def text(self) -> str:
        """All text, sections joined with their titles."""
        parts: list[str] = []
        for section in self.sections:
            parts.append(f"{section.title}\n\n{section.text}" if section.text else section.title)
        return "\n\n".join(parts)

    @property
    def section_titles(self) -> list[str]:
        return [s.title for s in self.sections]


def _looks_like_heading(line: str) -> bool:
    """Heuristic: short, title-caps-like line without terminal punctuation."""
    line = line.strip()
    if not line or len(line) > 90:
        return False
    if line.endswith((".", "!", "?", ",", ";", ":")):
        return False
    if len(line.split()) > _MAX_HEADING_WORDS:
        return False
    # Reject sentences: contain lowercase verbs patterns like " the ", " is "
    if re.search(r"\b(the|is|are|was|were|will|shall|must|of)\b", line, re.IGNORECASE) and not line[:1].isdigit():
        return False
    return bool(_HEADING_PATTERN.match(line))


def _group_paragraphs_into_sections(paragraphs: list[str], fallback_title: str) -> list[Section]:
    """Group a flat paragraph stream into sections using heading detection."""
    sections: list[Section] = []
    current: Section | None = None

    for paragraph in paragraphs:
        text = paragraph.strip()
        if not text:
            continue
        if _looks_like_heading(text):
            current = Section(title=text)
            sections.append(current)
        else:
            if current is None:
                current = Section(title=fallback_title)
                sections.append(current)
            current.paragraphs.append(text)

    if not sections:
        sections.append(Section(title=fallback_title))
    return sections


def _clean_paragraphs(raw_paragraphs: list[str]) -> list[str]:
    """Normalize whitespace; drop hyphenation artifacts and page markers."""
    cleaned: list[str] = []
    for raw in raw_paragraphs:
        text = raw.replace("\u00ad", "")  # soft hyphen
        text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line breaks
        text = re.sub(r"\s*\n\s*", " ", text)
        text = re.sub(r"[ \t]+", " ", text).strip()
        text = re.sub(r"^\s*(page\s+\d+(\s+of\s+\d+)?)\s*$", "", text, flags=re.IGNORECASE)
        if text:
            cleaned.append(text)
    return cleaned


def extract_pdf(path: Path) -> ExtractedDocument:
    """Extract text per PDF page (pypdf), then group into heading sections."""
    reader = PdfReader(str(path))
    extracted = ExtractedDocument()
    for page_number, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        paragraphs = _clean_paragraphs(raw.splitlines())
        page_sections = _group_paragraphs_into_sections(paragraphs, f"Page {page_number}")
        extracted.sections.extend(page_sections)
    return extracted


def extract_docx(path: Path) -> ExtractedDocument:
    """Extract text from DOCX using real paragraph styles for headings."""
    document = DocxDocument(str(path))
    extracted = ExtractedDocument()
    current: Section | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        is_heading_style = style_name.startswith("heading") or style_name in {"title", "subtitle"}
        if is_heading_style or (not paragraph.style.name and _looks_like_heading(text)):
            current = Section(title=text)
            extracted.sections.append(current)
        else:
            if current is None:
                current = Section(title="Document")
                extracted.sections.append(current)
            current.paragraphs.extend(_clean_paragraphs([text]))

    if not extracted.sections:
        extracted.sections.append(Section(title="Document"))
    return extracted


def extract_text(path: Path, file_type: str) -> ExtractedDocument:
    """Dispatch extraction by file type ('pdf' or 'docx')."""
    if file_type == "pdf":
        return extract_pdf(path)
    if file_type == "docx":
        return extract_docx(path)
    raise ValueError(f"Unsupported file type: {file_type!r}")