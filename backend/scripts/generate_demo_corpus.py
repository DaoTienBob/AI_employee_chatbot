"""Generate the demo corpus (one DOCX per knowledge topic).

Creates ``data/demo_corpus/*.docx`` plus ``manifest.json`` describing each
document's category, topic and suggested role visibility (see
docs/knowledge-taxonomy.md). The content is short placeholder text — replace
with the official documents for the real deployment (T26).

Usage (from the project root):

    .venv/bin/python backend/scripts/generate_demo_corpus.py

Re-running regenerates all files; wipe ``data/app.db`` first if you want a
clean slate on the server side.
"""

import json
from pathlib import Path

from docx import Document

ALL_ROLES = ["employee", "hr", "manager"]

CORPUS: dict[str, list[dict]] = {
    "Policy": [
        {
            "topic": "Nghỉ phép",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên chính thức được hưởng 12 ngày phép năm, tăng thêm theo thâm niên. "
                "Đơn xin nghỉ phép cần được quản lý trực tiếp phê duyệt trên hệ thống "
                "trước ít nhất 3 ngày làm việc.",
                "Nghỉ phép đột xuất do ốm đau cần thông báo cho quản lý trước 8 giờ sáng "
                "và bổ sung đơn trong vòng 2 ngày làm việc. Ngày phép không sử dụng trong "
                "năm sẽ được chuyển tối đa 5 ngày sang năm sau.",
            ],
        },
        {
            "topic": "Đi muộn/về sớm",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên đi muộn hoặc về sớm cần đăng ký trên hệ thống chấm công kèm lý do.",
                "Trường hợp đi muộn quá 3 lần trong một tháng sẽ được nhắc nhở; quá 5 lần "
                "sẽ ảnh hưởng đến đánh giá kỷ luật và mức thưởng tháng.",
            ],
        },
        {
            "topic": "Chấm công",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên chấm công qua ứng dụng nội bộ khi bắt đầu và kết thúc ngày làm việc.",
                "Bảng công được khóa vào ngày 25 hàng tháng. Nhân viên kiểm tra và gửi yêu cầu "
                "chỉnh sửa công cho bộ phận nhân sự trước thời điểm này nếu có sai lệch.",
            ],
        },
        {
            "topic": "WFH",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên được làm việc từ xa (WFH) tối đa 2 ngày mỗi tuần, cần được quản lý "
                "phê duyệt trước và đảm bảo tham gia họp trực tuyến đầy đủ.",
                "Trong thời gian WFH, nhân viên vẫn phải tuân thủ quy định bảo mật thông tin "
                "và sử dụng thiết bị do công ty cấp.",
            ],
        },
        {
            "topic": "OT",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Làm thêm giờ (OT) cần được phê duyệt bằng văn bản từ quản lý trước khi thực hiện.",
                "OT trong ngày thường được tính hệ số 150%, cuối tuần 200% và ngày lễ 300% "
                "theo quy định của bộ luật lao động.",
            ],
        },
        {
            "topic": "Chính sách phúc lợi",
            "roles": ["hr", "manager"],
            "paragraphs": [
                "Công ty thực hiện review lương và phúc lợi 2 lần mỗi năm dựa trên kết quả "
                "đánh giá hiệu suất và mặt bằng thị trường.",
                "Nhân viên được hưởng phúc lợi sức khỏe định kỳ hằng năm, quỹ đội nhóm, "
                "và các gói hỗ trợ theo cấp bậc. Chi tiết mức hưởng được phòng nhân sự "
                "thông báo riêng cho từng nhân viên.",
            ],
        },
        {
            "topic": "Bảo hiểm",
            "roles": ["hr", "manager"],
            "paragraphs": [
                "Công ty đóng bảo hiểm xã hội, y tế và thất nghiệp đầy đủ theo quy định pháp luật.",
                "Nhân viên chính thức được tham gia gói bảo hiểm sức khỏe cao cấp dành cho "
                "bản thân và có thể đăng ký cho người phụ thuộc với chi phí ưu đãi.",
            ],
        },
        {
            "topic": "Thưởng",
            "roles": ["hr", "manager"],
            "paragraphs": [
                "Nhân viên được hưởng lương tháng thứ 13 vào cuối năm dương lịch.",
                "Thưởng hiệu suất (KPI bonus) được đánh giá hằng năm dựa trên mục tiêu cá nhân "
                "và kết quả kinh doanh của công ty, chi trả vào quý 1 năm sau.",
            ],
        },
        {
            "topic": "Quy định thử việc",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Thời gian thử việc là 2 tháng đối với nhân viên chính thức và 1 tháng đối với "
                "vị trí thực tập sinh chuyển chính thức.",
                "Cuối kỳ thử việc, quản lý trực tiếp và phòng nhân sự đánh giá kết quả. "
                "Đạt yêu cầu sẽ ký hợp đồng chính thức; không đạt sẽ được thông báo trước 3 ngày.",
            ],
        },
        {
            "topic": "Quy trình nghỉ việc",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên muốn nghỉ việc cần gửi đơn trước ít nhất 30 ngày (45 ngày đối với "
                "vị trí quản lý) để đảm bảo thời gian bàn giao.",
                "Quy trình gồm: gửi đơn trên hệ thống, bàn giao công việc cho người tiếp nhận, "
                "hoàn tất clearance tài sản và tài khoản, rồi nhận quyết toán từ phòng nhân sự.",
            ],
        },
        {
            "topic": "Quy trình onboarding",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Trong ngày đầu tiên, nhân viên mới được phòng nhân sự hướng dẫn thủ tục, "
                "cấp tài khoản và thiết bị, sau đó được giới thiệu với các phòng ban.",
                "Mỗi nhân viên mới được phân công một buddy hỗ trợ trong 30 ngày đầu. "
                "Lộ trình onboarding 90 ngày gồm các mốc tuần 1, tuần 4 và tuần 12.",
            ],
        },
        {
            "topic": "Chính sách đào tạo",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Mỗi nhân viên có ngân sách đào tạo hằng năm để tham gia khóa học, "
                "hội thảo hoặc chứng chỉ liên quan đến công việc.",
                "Nhân viên đề xuất khóa học qua quản lý và phòng nhân sự phê duyệt. "
                "Sau khóa học, nhân viên chia sẻ kiến thức với nhóm trong buổi sharing nội bộ.",
            ],
        },
    ],
    "Operations": [
        {
            "topic": "Quy trình xin mua hàng",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên điền form xin mua hàng trên hệ thống, ghi rõ lý do, số lượng và "
                "ngân sách dự kiến, sau đó gửi quản lý trực tiếp phê duyệt.",
                "Đề xuất trên 5 triệu đồng cần thêm phê duyệt của trưởng phòng và bộ phận "
                "tài chính trước khi bộ phận hành chính thực hiện mua.",
            ],
        },
        {
            "topic": "Quy trình thanh toán",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Hóa đơn và chứng từ hợp lệ được gửi cho bộ phận kế toán, ghi rõ mã đề xuất "
                "hoặc mã dự án liên quan.",
                "Kế toán xử lý thanh toán theo chu kỳ hằng tuần vào thứ năm. "
                "Chứng từ không hợp lệ sẽ được trả lại để bổ sung trong 3 ngày làm việc.",
            ],
        },
        {
            "topic": "Công tác phí",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Công tác phí bao gồm vận chuyển, lưu trú và ăn uống theo chuẩn chi áp dụng "
                "cho từng cấp bậc. Nhân viên giữ toàn bộ hóa đơn trong chuyến đi.",
                "Trong vòng 7 ngày làm việc sau khi kết thúc công tác, nhân viên hoàn thành "
                "báo cáo chi phí đính kèm hóa đơn để được hoàn tiền.",
            ],
        },
        {
            "topic": "Mượn thiết bị",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên mượn thiết bị (laptop rời, máy chiếu, ổ cứng...) từ bộ phận IT "
                "qua form mượn thiết bị, ghi rõ thời gian và mục đích sử dụng.",
                "Người mượn chịu trách nhiệm bảo quản thiết bị và trả đúng hạn. "
                "Thiết bị hư hỏng do lỗi sử dụng sẽ được xử lý theo quy định tài sản.",
            ],
        },
        {
            "topic": "Đặt phòng họp",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Nhân viên đặt phòng họp qua hệ thống đặt lịch nội bộ, chọn khung giờ "
                "và phòng phù hợp với số lượng người tham dự.",
                "Vui lòng hủy phòng nếu không sử dụng để người khác có thể đặt. "
                "Họp ngoài giờ làm việc cần xác nhận từ bộ phận hành chính.",
            ],
        },
        {
            "topic": "Quy trình approval",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Mọi đề xuất (mua hàng, thanh toán, nghỉ phép, công tác phí) đều tuân theo "
                "ma trận phê duyệt theo cấp bậc và giá trị.",
                "Trình tự chuẩn: nhân viên, quản lý trực tiếp, trưởng phòng, rồi bộ phận "
                "liên quan (tài chính/nhân sự) khi vượt thẩm quyền. Mỗi bước được ghi lại "
                "trên hệ thống để truy vết.",
            ],
        },
        {
            "topic": "Quy định sử dụng tài sản",
            "roles": ALL_ROLES,
            "paragraphs": [
                "Tài sản được cấp phát (laptop, màn hình, điện thoại...) gắn mã định danh "
                "và được ghi nhận vào biên bản bàn giao cho từng nhân viên.",
                "Công ty kiểm kê tài sản định kỳ 2 lần mỗi năm. Khi nghỉ việc, toàn bộ "
                "tài sản phải được trả lại trong bước clearance trước ngày làm việc cuối.",
            ],
        },
    ],
    "Company tour": [
        {"topic": "Company handbook", "roles": ALL_ROLES, "paragraphs": [
            "Company handbook tổng hợp mọi quy định làm việc, quyền lợi và nghĩa vụ "
            "của nhân viên. Mỗi nhân viên ký nhận điện tử trong tuần đầu tiên onboarding.",
            "Handbook được cập nhật mỗi năm; khi có thay đổi quan trọng, phòng nhân sự "
            "sẽ gửi email thông báo kèm bản tóm tắt.",
        ]},
        {"topic": "Văn hóa", "roles": ALL_ROLES, "paragraphs": [
            "Văn hóa công ty đề cao sự cởi mở, tôn trọng và học hỏi lẫn nhau giữa "
            "các thành viên trong mọi cấp bậc.",
            "Chúng tôi khuyến khích trao đổi thẳng thắn trong họp, ghi nhận đóng góp "
            "cá nhân và tổ chức các buổi sharing hằng tháng.",
        ]},
        {"topic": "Giá trị công ty", "roles": ALL_ROLES, "paragraphs": [
            "Bốn giá trị cốt lõi: Trung thực, Chất lượng, Hợp tác và Tiến bộ.",
            "Các giá trị này được dùng làm căn cứ trong đánh giá hiệu suất hằng năm "
            "cùng với mục tiêu công việc (KPI/OKR).",
        ]},
        {"topic": "Cơ cấu tổ chức", "roles": ALL_ROLES, "paragraphs": [
            "Công ty gồm ban giám đốc và 4 khối: Khối kỹ thuật, Khối kinh doanh, "
            "Khối vận hành và Khối hỗ trợ (nhân sự, tài chính, hành chính, IT).",
            "Sơ đồ tổ chức chi tiết được công bố trên cổng thông tin nội bộ và cập nhật "
            "khi có thay đổi nhân sự cấp quản lý.",
        ]},
        {"topic": "Các phòng ban", "roles": ALL_ROLES, "paragraphs": [
            "Các phòng ban hiện tại: Kỹ thuật, Kinh doanh, Marketing, Vận hành, "
            "Nhân sự, Tài chính kế toán, Hành chính và IT.",
            "Mỗi phòng ban có một trưởng phòng phụ trách và một kênh liên hệ nội bộ riêng.",
        ]},
        {"topic": "Quy định nội bộ", "roles": ALL_ROLES, "paragraphs": [
            "Giờ làm việc từ 9:00 đến 18:00, nghỉ trưa 1 giờ. Trang phục business casual "
            "vào các ngày trong tuần.",
            "Khu vực làm việc giữ yên tĩnh phù hợp; họp nhóm vui lòng dùng phòng họp "
            "hoặc khu vực discussion. Điện thoại cá nhân vui lòng để chế độ im lặng.",
        ]},
        {"topic": "Event", "roles": ALL_ROLES, "paragraphs": [
            "Công ty tổ chức họp toàn công ty (all-hands) vào ngày đầu tiên của tháng, "
            "và sự kiện vui chơi theo mùa trong năm.",
            "Lịch sự kiện được công bố trên cổng thông tin nội bộ; đề xuất ý tưởng sự kiện "
            "gửi cho bộ phận hành chính hoặc đội văn hóa.",
        ]},
        {"topic": "Teambuilding", "roles": ALL_ROLES, "paragraphs": [
            "Mỗi nhóm được ngân sách team building hằng năm, tự chọn hình thức và "
            "thời gian phù hợp sau khi thống nhất với quản lý.",
            "Company trip hằng năm được tổ chức chung cho toàn công ty, thường rơi vào "
            "quý 2 hoặc quý 3.",
        ]},
        {"topic": "Sinh nhật", "roles": ALL_ROLES, "paragraphs": [
            "Sinh nhật nhân viên trong tháng được tổ chức chung vào chiều thứ sáu cuối "
            "tháng, có bánh và quà nhỏ từ công ty.",
            "Danh sách sinh nhật được bộ phận hành chính cập nhật trên bảng tin nội bộ.",
        ]},
        {"topic": "Chính sách employee benefit", "roles": ALL_ROLES, "paragraphs": [
            "Tổng quan phúc lợi chung cho mọi nhân viên: bảo hiểm sức khỏe, ngày phép, "
            "quỹ đội nhóm, chương trình chăm sóc sức khỏe định kỳ và ưu đãi đối tác.",
            "Xem chi tiết từng gói phúc lợi trong tài liệu chính sách phúc lợi do "
            "phòng nhân sự quản lý.",
        ]},
    ],
    "Contact": [
        {"topic": "Contact liên hệ", "roles": ALL_ROLES, "paragraphs": [
            "Nếu câu hỏi của bạn không tìm thấy thông tin trong các tài liệu, vui lòng "
            "liên hệ email phù hợp: hr@company.com cho chính sách nhân sự, phúc lợi "
            "và lương thưởng; it@company.com cho tài khoản, thiết bị và phần mềm; "
            "ops@company.com cho mua hàng, thanh toán và hành chính.",
            "Câu hỏi pháp lý hoặc hợp đồng gửi legal@company.com. Mọi yêu cầu khác "
            "gửi admin@company.com, chúng tôi sẽ chuyển đến bộ phận phụ trách.",
        ]},
    ],
}


def slugify(topic: str) -> str:
    """ASCII, filesystem-safe filename stem for a topic."""
    import unicodedata

    text = unicodedata.normalize("NFD", topic)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D")
    slug = "".join(ch if ch.isalnum() else "-" for ch in text).lower()
    return "-".join(part for part in slug.split("-") if part)


def main() -> None:
    out_dir = Path(__file__).resolve().parents[2] / "data" / "demo_corpus"
    out_dir.mkdir(parents=True, exist_ok=True)

    manifest = []
    for category, entries in CORPUS.items():
        for entry in entries:
            filename = f"{slugify(entry['topic'])}.docx"
            doc = Document()
            doc.add_heading(entry["topic"], level=0)
            doc.add_paragraph(f"Nhóm tài liệu: {category}")
            for paragraph in entry["paragraphs"]:
                doc.add_paragraph(paragraph)
            doc.save(out_dir / filename)
            manifest.append(
                {
                    "category": category,
                    "topic": entry["topic"],
                    "filename": filename,
                    "suggested_allowed_roles": entry["roles"],
                }
            )

    manifest_path = out_dir / "manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Wrote {len(manifest)} documents to {out_dir}")
    print(f"Manifest: {manifest_path}")


if __name__ == "__main__":
    main()
