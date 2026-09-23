"""Regression checks for the Phase 1–2 audit."""
import io
from pathlib import Path

import numpy as np
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import select

from backend.app.chunking import chunk_document
from backend.app.database import SessionLocal
from backend.app.main import app
from backend.app.models import Document, User
from backend.app.parsing import extract_docx, extract_pdf
from backend.app.vector_store import get_vector_store

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def docx_bytes(text):
    doc = DocxDocument()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(monkeypatch):
    with TestClient(app, raise_server_exceptions=False) as client:
        with SessionLocal() as db:
            existing = set(db.scalars(select(Document.id)))
        yield client
        monkeypatch.undo()
        with SessionLocal() as db:
            for document in db.scalars(select(Document)).all():
                if document.id not in existing:
                    get_vector_store().delete_document_chunks(f'DOC_{document.id:03d}')
                    Path(document.file_path).unlink(missing_ok=True)
                    db.delete(document)
            db.commit()


def login(client, role):
    response = client.post('/auth/login', json={
        'email': f'{role}@company.com', 'password': 'password123',
    })
    assert response.status_code == 200
    return {'Authorization': f"Bearer {response.json()['access_token']}"}


def upload(client, headers, text='Original policy content.', roles='hr', name='policy.docx'):
    response = client.post('/documents/upload', headers=headers,
        files={'file': (name, docx_bytes(text), MIME)}, data={'allowed_roles': roles})
    assert response.status_code == 200, response.text
    return response.json()['document']['id']


def test_docx_tables_keep_order_and_sections(tmp_path):
    doc = DocxDocument()
    doc.add_heading('Leave', level=1)
    doc.add_paragraph('Before table.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Twelve leave days.'
    doc.add_paragraph('After table.')
    doc.add_heading('Insurance', level=1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Dental coverage.'
    path = tmp_path / 'policy.docx'
    doc.save(path)
    result = extract_docx(path)
    assert result.sections[0].paragraphs == ['Before table.', 'Twelve leave days.', 'After table.']
    assert result.sections[1].text == 'Dental coverage.'


def test_table_only_docx_is_readable(tmp_path):
    doc = DocxDocument()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Twelve leave days.'
    path = tmp_path / 'table.docx'
    doc.save(path)
    assert extract_docx(path).body_text == 'Twelve leave days.'


def test_chunks_keep_section_attribution():
    chunks = chunk_document([('Leave', 'Twelve leave days.'), ('Insurance', 'Dental coverage.')],
        document_id='DOC_TEST', document_name='policy', allowed_roles=['hr'])
    assert [(c.section, c.text) for c in chunks] == [
        ('Leave', 'Twelve leave days.'), ('Insurance', 'Dental coverage.')]


@pytest.mark.parametrize('failure', ['parse', 'embedding', 'partial_index', 'commit'])
def test_failed_replace_preserves_all_stores(client, monkeypatch, failure):
    headers = login(client, 'admin')
    doc_id = upload(client, headers)
    store = get_vector_store()
    key = f'DOC_{doc_id:03d}'
    before = store.get_document_chunks(key)
    with SessionLocal() as db:
        original_path = Path(db.get(Document, doc_id).file_path)
    original_bytes = original_path.read_bytes()
    files_before = set(original_path.parent.iterdir())
    if failure in ('embedding', 'partial_index'):
        original_index = store.index_chunks
        def fail_index(chunks):
            if failure == 'partial_index':
                original_index(chunks)
            raise RuntimeError('Simulated embedding/index failure')
        monkeypatch.setattr(store, 'index_chunks', fail_index)
    elif failure == 'commit':
        def fail_commit(self):
            raise RuntimeError('Simulated SQL commit failure')
        monkeypatch.setattr(SessionLocal.class_, 'commit', fail_commit)
    payload = b'broken' if failure == 'parse' else docx_bytes('New policy. ' * 600)
    response = client.put(f'/documents/{doc_id}', headers=headers,
        files={'file': ('policy.docx', payload, MIME)}, data={'allowed_roles': 'employee'})
    assert response.status_code == (422 if failure == 'parse' else 500)
    assert original_path.read_bytes() == original_bytes
    assert set(original_path.parent.iterdir()) == files_before
    after = store.get_document_chunks(key)
    assert after['ids'] == before['ids']
    assert after['documents'] == before['documents']
    assert after['metadatas'] == before['metadatas']
    # Chroma normalizes embeddings on every upsert round-trip, so bit-exact
    # equality is not preserved; a tight tolerance confirms the original
    # vectors were restored (ids/documents/metadatas above are exact).
    assert np.allclose(after['embeddings'], before['embeddings'], atol=1e-5)
    with SessionLocal() as db:
        record = db.get(Document, doc_id)
        assert record.file_path == str(original_path)
        assert record.allowed_roles() == ['hr']
        assert record.content.endswith('Original policy content.')


def test_successful_same_name_replace(client):
    headers = login(client, 'admin')
    doc_id = upload(client, headers)
    with SessionLocal() as db:
        old_path = Path(db.get(Document, doc_id).file_path)
    response = client.put(f'/documents/{doc_id}', headers=headers,
        files={'file': ('policy.docx', docx_bytes('New policy.'), MIME)},
        data={'allowed_roles': 'manager'})
    assert response.status_code == 200
    assert not old_path.exists()
    with SessionLocal() as db:
        record = db.get(Document, doc_id)
        assert Path(record.file_path).exists()
        assert record.allowed_roles() == ['manager']
    chunks = get_vector_store().get_document_chunks(f'DOC_{doc_id:03d}')
    assert chunks['documents'] == ['New policy.']
    assert chunks['metadatas'][0]['allow_manager'] is True
    assert chunks['metadatas'][0]['allow_hr'] is False


def test_failed_upload_cleans_index_file_and_record(client, monkeypatch):
    headers = login(client, 'admin')
    store = get_vector_store()
    before = set(store._collection.get()['ids'])
    with SessionLocal() as db:
        records = list(db.scalars(select(Document.id)))
    from backend.app.config import get_settings
    folder = Path(get_settings().upload_dir)
    files = set(folder.iterdir())
    original = store.index_chunks
    def fail(chunks):
        original(chunks)
        raise RuntimeError('Partial indexing failure')
    monkeypatch.setattr(store, 'index_chunks', fail)
    response = client.post('/documents/upload', headers=headers,
        files={'file': ('failed.docx', docx_bytes('Must not persist.'), MIME)},
        data={'allowed_roles': 'employee'})
    assert response.status_code == 500
    assert set(store._collection.get()['ids']) == before
    assert set(folder.iterdir()) == files
    with SessionLocal() as db:
        assert list(db.scalars(select(Document.id))) == records


@pytest.mark.parametrize('role', ['employee', 'hr', 'manager'])
def test_authenticated_search_role_isolation(client, role):
    admin = login(client, 'admin')
    ids = {}
    for allowed in ['employee', 'hr', 'manager']:
        ids[allowed] = upload(client, admin, f'Audit unique {allowed} policy.', roles=allowed)
    response = client.get('/documents/search', headers=login(client, role),
        params={'query': f'Audit unique {role} policy.', 'role': 'hr'})
    assert response.status_code == 200
    hits = response.json()
    assert hits
    assert all(hit['metadata'][f'allow_{role}'] for hit in hits)
    returned = {hit['metadata']['document_id'] for hit in hits}
    assert f'DOC_{ids[role]:03d}' in returned
    assert not returned.intersection(f'DOC_{id:03d}' for other, id in ids.items() if other != role)


def test_search_requires_session_and_uses_current_role(client):
    assert client.get('/documents/search', params={'query': 'policy'}).status_code == 401
    headers = login(client, 'hr')
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == 'hr@company.com'))
        user.role = 'employee'
        db.commit()
    try:
        response = client.get('/documents/search', headers=headers, params={'query': 'policy'})
        assert response.status_code == 200
        assert all(hit['metadata']['allow_employee'] for hit in response.json())
    finally:
        with SessionLocal() as db:
            user = db.scalar(select(User).where(User.email == 'hr@company.com'))
            user.role = 'hr'
            db.commit()


def test_pdf_extraction_and_chunking(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
        DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (Employees receive twelve leave days.) Tj ET')
    page[NameObject('/Contents')] = stream
    path = tmp_path / 'policy.pdf'
    writer.write(path)
    extracted = extract_pdf(path)
    assert 'twelve leave days' in extracted.body_text
    chunks = chunk_document([(s.title, s.text) for s in extracted.sections],
        document_id='DOC_PDF', document_name='policy.pdf', allowed_roles=['employee'])
    assert chunks and chunks[0].section == 'Page 1'


@pytest.mark.parametrize('case, expected', [
    ('non_admin', 403), ('extension', 400), ('mime', 400),
    ('size', 413), ('roles', 400), ('empty', 400), ('corrupt', 422),
])
def test_upload_validation(client, monkeypatch, case, expected):
    from backend.app.config import get_settings
    headers = login(client, 'employee' if case == 'non_admin' else 'admin')
    name = 'policy.txt' if case == 'extension' else 'policy.docx'
    mime = 'application/pdf' if case == 'mime' else MIME
    payload = b'' if case == 'empty' else b'broken' if case == 'corrupt' else docx_bytes('Policy content.')
    roles = 'unknown' if case == 'roles' else 'employee'
    if case == 'size':
        monkeypatch.setattr(get_settings(), 'max_upload_size_mb', 0)
    response = client.post('/documents/upload', headers=headers,
        files={'file': (name, payload, mime)}, data={'allowed_roles': roles})
    assert response.status_code == expected, response.text


def test_health_and_login(client):
    assert client.get('/health').json()['status'] == 'ok'
    assert client.post('/auth/login', json={
        'email': 'employee@company.com', 'password': 'wrong',
    }).status_code == 401
    for role in ['employee', 'hr', 'manager']:
        response = client.get('/auth/me', headers=login(client, role))
        assert response.status_code == 200
        assert response.json()['role'] == role
        assert 'hashed_password' not in response.json()


def test_invalid_jwt_subject_returns_401(client):
    import jwt
    from backend.app.config import get_settings

    # Token with non-integer subject string
    payload = {"sub": "not-an-int", "role": "employee"}
    bad_token = jwt.encode(payload, get_settings().secret_key, algorithm="HS256")
    response = client.get("/auth/me", headers={"Authorization": f"Bearer {bad_token}"})
    assert response.status_code == 401
    assert response.json()["detail"] == "Could not validate credentials"


def test_answer_question_creates_single_conversation():
    from backend.app.database import SessionLocal
    from backend.app.models import Conversation, User
    from backend.app.rag import answer_question

    with SessionLocal() as db:
        user = db.query(User).filter_by(email="employee@company.com").first()
        conv_count_before = db.query(Conversation).filter_by(user_id=user.id).count()

        # Run answer_question with conversation_id=None
        resp = answer_question("What are the working hours?", user, db, conversation_id=None)
        assert resp.conversation_id is not None

        conv_count_after = db.query(Conversation).filter_by(user_id=user.id).count()
        # Exactly ONE new conversation must be created
        assert conv_count_after == conv_count_before + 1

